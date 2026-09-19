import base64
import hashlib
import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock

import pytest
from PIL import Image, ImageDraw

from src.config import settings
from src.db.metadata import MetadataStore
from src.figures import service, storage
from src.ingestion.figure_extract import ImageRegion, FigureRecord, ImageKind


def png():
    image = Image.new("RGB", (900, 500), "white")
    draw = ImageDraw.Draw(image)
    draw.ellipse((80, 80, 450, 450), outline="blue", width=4)
    draw.ellipse((300, 80, 670, 450), outline="red", width=4)
    draw.text((180, 240), "BRANCH A", fill="black")
    draw.text((500, 240), "BRANCH B", fill="black")
    result = io.BytesIO()
    image.save(result, "PNG")
    return result.getvalue()


@pytest.fixture
def figures_root(tmp_path, monkeypatch):
    monkeypatch.setattr(storage, "FIGURE_ROOT", tmp_path / "figures")
    monkeypatch.setattr(settings, "figure_store_enabled", True)
    return tmp_path / "figures"


def staged_record(tmp_path, figure_id="p1-fig-001"):
    source = tmp_path / "worker" / "figures"
    region = ImageRegion(0, 0, png(), 900, 500, figure_id=figure_id, caption="Branch overlap")
    with storage.extraction_assets(source):
        assets = storage.save_region(region)
    record = FigureRecord(figure_id, "Branch A and B overlap in a Venn diagram", "other", page=0, caption=region.caption, assets=assets)
    return storage.FigureStore().handoff(source), record


async def populated(tmp_path):
    ms = MetadataStore("sqlite+aiosqlite:///" + str(tmp_path / "metadata.db"))
    await ms.init()
    staging, record = staged_record(tmp_path)
    details = storage.FigureStore().publish("doc-1", [record], staging)
    await ms.put_figures("doc-1", details)
    await ms.add_document("doc-1", "topology.pdf", "pdf", ["network"], 1, "tester", content_hash="source-v1")
    storage.FigureStore().discard(staging)
    return ms, record


def test_png_handoff_hash_dedup_and_restart(tmp_path, figures_root):
    staging, record = staged_record(tmp_path)
    record2 = FigureRecord("p2-fig-001", "Same image on page two", "other", page=1, assets=record.assets)
    details = storage.FigureStore().publish("doc-1", [record, record2], staging)
    storage.FigureStore().discard(staging)
    assert len(details) == 2
    keys = {a["key"] for a in record.assets.values()}
    assert len(list((figures_root / "doc-1").glob("*.png"))) == len(keys)
    for asset in record.assets.values():
        path = storage.FigureStore().asset_path("doc-1", asset["key"])
        assert hashlib.sha256(path.read_bytes()).hexdigest() == asset["sha256"]
        assert Image.open(path).size == (900, 500)
    assert not Path(staging).exists()


def test_handoff_rejects_symlinks_and_hash_tampering(tmp_path, figures_root):
    directory = tmp_path / "bad"
    directory.mkdir()
    outside = tmp_path / "outside.png"
    outside.write_bytes(png())
    link = directory / ("a" * 64 + ".png")
    link.symlink_to(outside)
    with pytest.raises(ValueError):
        storage.FigureStore().handoff(directory)
    link.unlink()
    link.write_bytes(png())
    with pytest.raises(ValueError, match="hash"):
        storage.FigureStore().handoff(directory)
    assert not list((figures_root / ".staging").iterdir())
    with pytest.raises(ValueError):
        storage.FigureStore().asset_path("../escape", "a" * 64 + ".png")


def test_retains_image_when_vision_fails(tmp_path, figures_root, monkeypatch):
    import src.ingestion.figure_extract as f
    monkeypatch.setattr(settings, "figure_ocr_first", False)
    monkeypatch.setattr(f, "classify_region", lambda _: ImageKind.NETWORK)
    monkeypatch.setattr(f, "run_strategy", MagicMock(side_effect=RuntimeError("vision unavailable")))
    region = ImageRegion(0, 0, png(), 900, 500, figure_id="f1", caption="Branch Venn")
    with storage.extraction_assets(tmp_path / "worker"):
        result = f.process_image_regions([region])
    assert result.figure_records[0].assets["preview"]
    assert result.figure_records[0].analysis_status == "unavailable"
    assert "Branch Venn" in result.figure_records[0].retrieval_text()


def test_vector_venn_on_text_heavy_pdf_is_rendered(tmp_path, monkeypatch):
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    from src.ingestion.figure_extract import extract_image_regions
    path = tmp_path / "venn.pdf"
    writer = PdfWriter()
    page = writer.add_blank_page(612, 792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    # Two closed curved paths overlap, alongside substantial digital text.
    stream.set_data(b"BT /F1 12 Tf 50 750 Td (Branch overlap Venn diagram with enough digital text to exceed the sparse threshold.) Tj ET "
                    b"100 400 m 100 250 350 250 350 400 c 350 550 100 550 100 400 c S "
                    b"250 400 m 250 250 500 250 500 400 c 500 550 250 550 250 400 c S")
    page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as handle:
        writer.write(handle)
    monkeypatch.setattr(settings, "figure_render_vector_pages", True)
    regions = extract_image_regions(path)
    assert any(r.source == "page_render" and r.image_bytes.startswith(storage.PNG) for r in regions)
    monkeypatch.setattr(settings, "figure_render_vector_pages", False)
    assert extract_image_regions(path) == []


@pytest.mark.asyncio
async def test_acl_rechecked_for_bytes_cache_and_deletion(tmp_path, figures_root):
    ms, record = await populated(tmp_path)
    try:
        raw, ref = await service.image_bytes("doc-1", record.figure_id, ["network"], ms)
        assert raw.startswith(storage.PNG)
        assert ref["page"] == 1
        citation = {"doc_id": "doc-1", "figure_id": record.figure_id, "evidence_id": "E1"}
        assert await service.answer_images("Topology", [citation], ["network"], ms)
        with pytest.raises(FileNotFoundError):
            await service.image_bytes("doc-1", record.figure_id, ["finance"], ms)
        await ms.update_document("doc-1", acl_groups=["finance"])
        assert await service.answer_images("Topology", [citation], ["network"], ms) == []
        await ms.delete_document("doc-1")
        assert await ms.get_figure("doc-1", record.figure_id) is None
        assert not (figures_root / "doc-1").exists()
        with pytest.raises(FileNotFoundError):
            await service.image_bytes("doc-1", record.figure_id, ["ALL"], ms)
    finally:
        await ms.engine.dispose()


@pytest.mark.asyncio
async def test_pending_assets_are_not_retrievable(tmp_path, figures_root):
    ms = MetadataStore("sqlite+aiosqlite:///" + str(tmp_path / "metadata.db"))
    await ms.init()
    staging, record = staged_record(tmp_path)
    try:
        records = storage.FigureStore().publish("pending", [record], staging)
        await ms.put_figures("pending", records)
        with pytest.raises(FileNotFoundError):
            await service.image_bytes("pending", record.figure_id, ["ALL"], ms)
        await storage.FigureStore().reconcile(ms)
        assert not (figures_root / "pending").exists()
        assert not Path(staging).exists()
    finally:
        await ms.engine.dispose()


@pytest.mark.asyncio
async def test_native_mcp_image_roundtrip(tmp_path, figures_root):
    from fastmcp import FastMCP, Client
    ms, record = await populated(tmp_path)
    try:
        doc, fig = await service.authorized_figure("doc-1", record.figure_id, ["network"], ms)
        ref = service.reference(doc, fig)
        server = FastMCP("diagram-test")
        @server.tool()
        async def diagram():
            return await service.mcp_result({"answer": "Source Venn diagram", "images": [ref]}, ["network"], ms)
        async with Client(server) as client:
            result = await client.call_tool("diagram")
        assert [b.type for b in result.content] == ["text", "image"]
        assert base64.b64decode(result.content[1].data).startswith(storage.PNG)
        assert "base64" not in result.content[0].text
        metadata = json.loads(result.content[0].text)
        assert metadata["images"][0]["figure_id"] == record.figure_id
        assert result.structured_content["answer"] == "Source Venn diagram"
        await ms.update_document("doc-1", acl_groups=["finance"])
        result = await service.mcp_result({"images": [ref]}, ["network"], ms)
        assert len(result.content) == 1
        assert result.structured_content["images"] == []
        assert result.structured_content["warnings"]
    finally:
        await ms.engine.dispose()


def test_figure_filter_precedes_top_k(tmp_path, monkeypatch):
    from src.retrieval.vector_store import VectorStore
    from src.retrieval.models import ChunkMetadata
    monkeypatch.setattr(settings, "lancedb_path", str(tmp_path / "vectors"))
    monkeypatch.setattr(settings, "embedding_dimension", 3)
    store = VectorStore()
    text_meta = [ChunkMetadata(doc_id="d", filename="d.pdf", doc_type="pdf", chunk_index=i, start_char=0, acl_groups=["network"]) for i in range(25)]
    figure = text_meta[0].model_copy(update={"content_type": "figure", "figure_id": "f1", "figure_kind": "network", "chunk_index": 26})
    store.upsert(["Topology prose"] * 25 + ["Branch topology diagram"], [[1., 0., 0.]] * 25 + [[.8, .2, 0.]], text_meta + [figure])
    results = store.search_figures([1., 0., 0.], "topology", ["network"], ["d"], top_k=1)
    assert len(results) == 1
    assert results[0].metadata.figure_id == "f1"
    assert store.search_figures([1., 0., 0.], "topology", ["finance"], ["d"], top_k=1) == []


@pytest.mark.asyncio
async def test_asset_only_changes_invalidate_query_scope(tmp_path, figures_root, monkeypatch):
    from src.retrieval.query_scope import resolve_query_scope
    monkeypatch.setattr(settings, "lancedb_path", str(tmp_path / "no-index"))
    ms, record = await populated(tmp_path)
    try:
        before = await resolve_query_scope(["network"], ms)
        await ms.put_figures("doc-1", [])
        after = await resolve_query_scope(["network"], ms)
        assert before.doc_ids == after.doc_ids
        assert before.revision != after.revision
    finally:
        await ms.engine.dispose()


def test_policy_and_resource_limits(tmp_path, figures_root, monkeypatch):
    monkeypatch.setattr(settings, "answer_images", "auto")
    monkeypatch.setattr(settings, "answer_max_images", 2)
    assert service.image_policy("topology", {"config": {"images": "off"}}) == 0
    assert service.image_policy("deployment", {"config": {"images": "requested"}}) == 0
    assert service.image_policy("show topology", {"config": {"images": "requested", "max_images": 1}}) == 1
    monkeypatch.setattr(settings, "figure_max_pixels", 10000)
    with storage.extraction_assets(tmp_path / "worker"):
        with pytest.raises(ValueError, match="pixel"):
            storage.save_region(ImageRegion(0, 0, png(), 900, 500))


@pytest.mark.asyncio
async def test_actual_worker_transfers_png_before_cleanup(tmp_path, figures_root, monkeypatch):
    import sys
    from docx import Document
    from docx.shared import Inches
    from src.ingestion import extraction_worker
    from src.ingestion.isolation import extract_in_worker
    from src.ingestion.prepared_index import index_prepared
    source = tmp_path / "network.docx"
    image = tmp_path / "figure.png"
    image.write_bytes(png())
    doc = Document()
    doc.add_heading("Branch overlap", 1)
    doc.add_picture(str(image), width=Inches(4))
    doc.save(str(source))
    monkeypatch.setattr(settings, "extraction_work_dir", str(tmp_path / "extraction"))
    monkeypatch.setattr(settings, "figure_ocr_first", False)
    real = extraction_worker.run_task
    async def run(task_dir):
        script = (
            "from pathlib import Path; import src.ingestion.figure_extract as f; "
            "from src.ingestion.extraction_worker import child; "
            "f.classify_region=lambda r: f.ImageKind.NETWORK; "
            "f.run_strategy=lambda r: ([], [f.ProseBlock('Branch network diagram', r.page)]); "
            f"child(Path({str(task_dir)!r}))"
        )
        await real(task_dir, command=[sys.executable, "-c", script])
    monkeypatch.setattr(extraction_worker, "run_task", run)
    result = await extract_in_worker(source, "network.docx")
    assert not list(Path(settings.extraction_work_dir).iterdir())
    assert Path(result.figure_staging).is_dir()
    assert result.office.figures[0].assets
    serialized = json.dumps(__import__('src.ingestion.prepared', fromlist=['encode_prepared']).encode_prepared(result))
    assert "base64" not in serialized
    ms = MetadataStore("sqlite+aiosqlite:///" + str(tmp_path / "worker.db"))
    await ms.init()
    try:
        await index_prepared(result, "worker-doc", ["network"], "", MagicMock(), ms)
        await ms.add_document("worker-doc", "network.docx", "docx", ["network"], 1, "test")
        storage.FigureStore().discard(result.figure_staging)
        raw, _ = await service.image_bytes("worker-doc", result.office.figures[0].figure_id, ["network"], ms)
        assert raw.startswith(storage.PNG)
    finally:
        await ms.engine.dispose()


@pytest.mark.asyncio
async def test_backfill_preserves_document_and_rolls_back_failure(tmp_path, figures_root, monkeypatch):
    from src.figures.backfill import backfill_figures
    from src.ingestion.prepared import PreparedDocument
    from src.ingestion.parser import ParsedDocument
    from src.ingestion.figure_extract import OfficeFigureResult
    from src.retrieval.vector_store import VectorStore
    from src.retrieval.models import ChunkMetadata
    monkeypatch.setattr(settings, "lancedb_path", str(tmp_path / "lance"))
    monkeypatch.setattr(settings, "embedding_dimension", 3)
    ms = MetadataStore("sqlite+aiosqlite:///" + str(tmp_path / "backfill.db"))
    await ms.init()
    source = tmp_path / "original.docx"
    source.write_bytes(b"same original bytes")
    checksum = hashlib.sha256(source.read_bytes()).hexdigest()
    await ms.add_document("doc-1", source.name, "docx", ["network"], 2, "owner", content_hash=checksum)
    vectors = VectorStore()
    text_meta = ChunkMetadata(doc_id="doc-1", filename=source.name, doc_type="docx", chunk_index=0, start_char=0, acl_groups=["network"])
    old_fig = text_meta.model_copy(update={"chunk_index": 1, "figure_id": "f1", "content_type": "figure"})
    vectors.upsert(["preserved text", "old diagram description"], [[1.,0.,0.]] * 2, [text_meta, old_fig])
    async def extract(*args):
        staging, record = staged_record(tmp_path, "f1")
        return PreparedDocument(ParsedDocument(source.name, "docx", "preserved text"), office=OfficeFigureResult("text", figures=[record]), figure_staging=staging)
    monkeypatch.setattr("src.ingestion.isolation.extract_in_worker", extract)
    monkeypatch.setattr("src.ingestion.embedder.embed_texts", lambda texts: [[1.,0.,0.]] * len(texts))
    real_put = ms.put_figures
    monkeypatch.setattr(ms, "put_figures", AsyncMock(side_effect=RuntimeError("db full")))
    try:
        with pytest.raises(RuntimeError, match="db full"):
            await backfill_figures("doc-1", source, ms, vectors)
        chunks = vectors.get_chunks_by_doc("doc-1")
        assert len(chunks) == 2
        assert {c.text for c in chunks} == {"preserved text", "old diagram description"}
        assert not list((figures_root / ".staging").iterdir())
        assert not (figures_root / "doc-1").exists()
        monkeypatch.setattr(ms, "put_figures", real_put)
        result = await backfill_figures("doc-1", source, ms, vectors)
        assert result["stored"] == 1
        assert (await ms.get_document("doc-1")).acl_groups == ["network"]
        assert (await ms.get_document("doc-1")).uploaded_by == "owner"
        chunks = vectors.get_chunks_by_doc("doc-1")
        assert len(chunks) == 2
        assert "preserved text" in [c.text for c in chunks]
        assert await ms.get_figure("doc-1", "stored-f1")
        with pytest.raises(ValueError, match="already has"):
            await backfill_figures("doc-1", source, ms, vectors)
        source.write_bytes(b"different")
        with pytest.raises(ValueError, match="differs"):
            await backfill_figures("doc-1", source, ms, vectors)
    finally:
        await ms.engine.dispose()


@pytest.mark.asyncio
async def test_http_authorization_and_admin_page(tmp_path, figures_root, monkeypatch):
    import httpx
    from src.main import app
    from src.auth.jwt import create_token
    from src.auth.models import UserContext
    from src.api import routes_ingest
    from src.admin import routes as admin
    ms, record = await populated(tmp_path)
    monkeypatch.setattr(routes_ingest, "_metadata_store", ms)
    path = f"/api/v1/documents/doc-1/figures/{record.figure_id}/content"
    try:
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test") as client:
            assert (await client.get(path)).status_code == 403
            # Native endpoint requires a user token as well as a valid app key.
            assert (await client.get(path, headers={"X-API-Key": "test-key-1"})).status_code == 401
            token = create_token("tester", ["network"])
            response = await client.get(path, headers={"X-API-Key": "test-key-1", "Authorization": "Bearer " + token})
            assert response.status_code == 200
            assert response.content.startswith(storage.PNG)
            assert response.headers["cache-control"] == "private, no-store"
            denied = create_token("finance", ["finance"])
            assert (await client.get(path, headers={"X-API-Key": "test-key-1", "Authorization": "Bearer " + denied})).status_code == 404
            assert (await client.get('/admin/api/figure-documents/doc-1/figures/' + record.figure_id + '/content')).status_code == 401
            monkeypatch.setattr(admin, "_is_authenticated", lambda _: True)
            page = await client.get('/admin/diagrams')
            assert page.status_code == 200
            assert 'diagram-search' in page.text
            response = await client.get('/admin/api/figure-documents/doc-1/figures/' + record.figure_id + '/content')
            assert response.content.startswith(storage.PNG)
    finally:
        await ms.engine.dispose()


@pytest.mark.asyncio
async def test_mcp_payload_budget_does_not_claim_omitted_image(tmp_path, figures_root, monkeypatch):
    ms, record = await populated(tmp_path)
    try:
        doc, fig = await service.authorized_figure("doc-1", record.figure_id, ["network"], ms)
        ref = service.reference(doc, fig)
        monkeypatch.setattr(settings, "figure_mcp_max_mb", 0)
        result = await service.mcp_result({"answer": "text survives", "images": [ref]}, ["network"], ms)
        assert len(result.content) == 1
        assert result.structured_content["images"] == []
        assert "limit" in result.structured_content["warnings"][0]
    finally:
        await ms.engine.dispose()


@pytest.mark.asyncio
async def test_backup_restores_png_and_committed_metadata_without_staging(tmp_path, monkeypatch):
    import asyncio
    import sqlite3
    import tarfile
    from src.admin import routes
    from src.ingestion.queue import ingest_queue

    monkeypatch.chdir(tmp_path)
    (tmp_path / 'data/figures/doc-1').mkdir(parents=True)
    (tmp_path / 'data/figures/doc-1/test.png').write_bytes(png())
    (tmp_path / 'data/figures/.staging/pending').mkdir(parents=True)
    (tmp_path / 'data/figures/.staging/pending/partial.png').write_bytes(b'partial')
    (tmp_path / 'data/extraction/job').mkdir(parents=True)
    (tmp_path / 'data/extraction/job/source.pdf').write_bytes(b'temporary')
    source = sqlite3.connect('data/metadata.db')
    source.execute('PRAGMA journal_mode=WAL')
    source.execute('CREATE TABLE figures (id TEXT)')
    source.execute("INSERT INTO figures VALUES ('doc-1')")
    source.commit()
    monkeypatch.setattr(routes, '_backup_status', {'state': 'idle', 'message': ''})
    monkeypatch.setattr(ingest_queue, 'has_active_jobs', lambda: False)
    monkeypatch.setattr(storage, '_active_writes', 0)
    try:
        result = await routes.create_backup()
        assert result.status_code == 200
        await asyncio.sleep(0)
        assert routes._backup_status['state'] == 'done', routes._backup_status
        archive_path = next((tmp_path / 'backups').glob('*.tar.gz'))
        restored = tmp_path / 'restored'
        with tarfile.open(archive_path) as archive:
            names = archive.getnames()
            assert not any('.staging' in n or 'extraction' in n or n.endswith(('-wal', '-shm')) for n in names)
            for name in ('data/figures/doc-1/test.png', 'data/metadata.db'):
                destination = restored / name
                destination.parent.mkdir(parents=True, exist_ok=True)
                destination.write_bytes(archive.extractfile(name).read())
        assert (restored / 'data/figures/doc-1/test.png').read_bytes() == png()
        with sqlite3.connect(restored / 'data/metadata.db') as restored_db:
            assert restored_db.execute('SELECT id FROM figures').fetchall() == [('doc-1',)]
    finally:
        source.close()


@pytest.mark.asyncio
async def test_cancelled_publication_finishes_before_cleanup(tmp_path, figures_root, monkeypatch):
    import asyncio
    import threading
    store = storage.FigureStore()
    entered, release = threading.Event(), threading.Event()
    finished = []
    def copying(*args):
        entered.set()
        release.wait(timeout=5)
        finished.append(True)
        return []
    monkeypatch.setattr(store, 'publish', copying)
    task = asyncio.create_task(store.publish_async('doc-1', [], 'unused'))
    await asyncio.to_thread(entered.wait, 5)
    task.cancel()
    await asyncio.sleep(0)
    assert not task.done()
    release.set()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert finished == [True]


@pytest.mark.asyncio
async def test_pdf_table_failure_preserves_figure_publication(tmp_path, figures_root, monkeypatch):
    from src.ingestion import prepared_index
    staging, record = staged_record(tmp_path)
    prepared = SimpleNamespace(parsed=SimpleNamespace(doc_type='pdf', filename='source.pdf'),
        warnings=[], figure_staging=staging,
        pdf=SimpleNamespace(table_grids=[], prose_blocks=[], figure_records=[record]))
    monkeypatch.setattr(prepared_index, 'ingest_grids', AsyncMock(side_effect=RuntimeError('table failed')))
    ms = SimpleNamespace(put_figures=AsyncMock())
    _, _, records = await prepared_index.index_prepared(prepared, 'doc-1', ['network'], '', None, ms)
    assert records == [record]
    ms.put_figures.assert_awaited_once()
    assert (figures_root / 'doc-1' / record.assets['preview']['key']).exists()
